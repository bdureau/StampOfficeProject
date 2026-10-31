"""
Utility Helper Module
---------------------
Contains photo name resolution, case-insensitive file lookups, and Excel export logic.
The Excel headers are written in the language active at the time of the export.

Author: Boris du Reau
Date: 2026-10-31
"""

from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
# Translation function of the current language (follows a change of language)
from SCAppLanguage import _


def get_photo_name(stamp_number: str, stamp_type: str, stamp_sub_nbr) -> str:
    """Generate relative photo image file name based on stamp category prefix."""
    sub_nbr = ""
    if stamp_sub_nbr is not None and str(stamp_sub_nbr).strip() not in ["", "None"]:
        sub_nbr = "-" + str(stamp_sub_nbr)

    stamp_nbr_str = str(stamp_number)
    type_prefix = stamp_nbr_str[0:2]

    # Airmail
    if type_prefix == "A ":
        return "T01-030-" + stamp_nbr_str[2:] + sub_nbr + ".jpg"
    # Miniature sheet
    elif type_prefix == "BF":
        return "T01-080-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Official / Service
    elif type_prefix == "S ":
        return "T01-170-" + stamp_nbr_str[2:] + sub_nbr + ".jpg"
    # Booklet sheet
    elif type_prefix == "BC":
        return "T01-000-BC" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Strip
    elif type_prefix == "B ":
        return "T01-000-B" + stamp_nbr_str[2:] + sub_nbr + ".jpg"
    # Pair
    elif type_prefix == "P ":
        return "T01-000-P" + stamp_nbr_str[2:] + sub_nbr + ".jpg"
    # Pre-cancelled
    elif type_prefix == "PR":
        return "T01-420-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Military franchise
    elif type_prefix == "FM":
        return "T01-190-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Postage due / Tax
    elif type_prefix == "TA":
        return "T01-270-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Red Cross
    elif type_prefix == "CR":
        return "T01-565-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # War stamps
    elif type_prefix == "GU":
        return "T01-150-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Alsace Lorraine
    elif type_prefix == "AL":
        return "T01-110-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Newspapers
    elif type_prefix == "JN":
        return "T01-160-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Small parcel
    elif type_prefix == "PC":
        return "T01-120-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Postal parcel
    elif type_prefix == "CP":
        return "T01-130-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Fiscal postal
    elif type_prefix == "FP":
        return "T01-140-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # CFA
    elif type_prefix == "CF":
        return "T01-180-CF-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Colonies
    elif type_prefix == "CO":
        return "T01-200-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Telegraph
    elif type_prefix == "TL":
        return "T01-210-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Military telegraph
    elif type_prefix == "MT":
        return "T01-211-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Telephone
    elif type_prefix == "TE":
        return "T01-215-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Strike
    elif type_prefix == "GR":
        return "T01-220-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Montenegro postal
    elif type_prefix == "PM":
        return "T01-230-" + stamp_nbr_str[3:] + sub_nbr + ".jpg"
    # Fictitious
    elif type_prefix == "F ":
        return "T01-240-" + stamp_nbr_str[2:] + sub_nbr + ".jpg"
    elif type_prefix == "F-":
        return "T01-000-F" + stamp_nbr_str[2:] + sub_nbr + ".jpg"
    # Charity
    elif type_prefix == "BI":
        return "T01-250-" + stamp_nbr_str[2:] + sub_nbr + ".jpg"
    # Money order
    elif type_prefix == "MA":
        return "T01-251-" + stamp_nbr_str[2:] + sub_nbr + ".jpg"
    # Vending machine / Dispenser
    elif type_prefix == "DI":
        return "T01-260-" + stamp_nbr_str[2:] + sub_nbr + ".jpg"
    else:
        stamp_type_str = str(stamp_type) if stamp_type else ""
        if stamp_type_str == "COURS INSTRUCTION":
            return "T01-252-" + stamp_nbr_str + sub_nbr + ".jpg"
        elif stamp_type_str == "Préoblitéré COURS INSTRUCTION":
            return "T01-253-" + stamp_nbr_str + sub_nbr + ".jpg"
        elif stamp_type_str == "RECOUVREMENT COURS INSTRUCTION":
            return "T01-254-" + stamp_nbr_str + sub_nbr + ".jpg"
        elif stamp_type_str.strip() == "TAXES COURS INSTRUCTION":
            return "T01-255-" + stamp_nbr_str + sub_nbr + ".jpg"
        elif stamp_type_str.strip() == "POSTE AERIENNE MILITAIRE":
            return "T01-256-" + stamp_nbr_str + sub_nbr + ".jpg"
        else:
            return "T01-000-" + stamp_nbr_str + sub_nbr + ".jpg"


def find_file_case_insensitive(folder_path: Path, filename: str) -> Path:
    """Search for a file in a directory ignoring case differences."""
    if not folder_path.is_dir():
        return None

    # Fast direct check (Windows or exact match on Linux)
    exact_path = folder_path / filename
    if exact_path.exists():
        return exact_path

    # Fallback case-insensitive search for Linux systems
    target_lower = filename.lower()
    for item in folder_path.iterdir():
        if item.is_file() and item.name.lower() == target_lower:
            return item

    return None


def export_mancolist_to_excel(file_path: str, mancolist: list):
    """Write missing stamp records into a formatted Excel (.xlsx) file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = _("Mancolist")

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1D6F42", end_color="1D6F42", fill_type="solid")

    headers = [
        _("Yvert N°"),
        _("Year"),
        _("Type"),
        _("Value / Color"),
        _("Description"),
        _("Status"),
        _("Notes")
    ]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_data in mancolist:
        ws.append(list(row_data))

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(file_path)