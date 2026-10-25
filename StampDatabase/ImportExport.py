"""
ImportExport Module
-------------------
Handles CSV and XLSX import/export logic for stamp catalogues.
"""

import csv
from PyQt6.QtWidgets import (
    QDialog, QVBoxLayout, QLabel, QRadioButton, QDialogButtonBox,
    QFileDialog, QMessageBox
)

import gettext
from SDAppLanguage import get_language

gettext.find("StampDatabase")
translate = gettext.translation(
    'StampDatabase', localedir='locale', languages=[get_language()], fallback=True
)
translate.install()
_ = translate.gettext

HEADERS = [
    "key", "nbr", "year", "type", "valuecolor",
    "stampDescription1", "stampDescription", "width",
    "height", "sequence", "ascii_seq", "sub_nbr"
]


class CsvDelimiterDialog(QDialog):
    """Dialog to select delimiter for CSV Import/Export (Comma selected by default)."""

    def __init__(self, parent=None, title="CSV Delimiter"):
        super().__init__(parent)
        self.setWindowTitle(_(title))
        self.resize(300, 120)

        layout = QVBoxLayout(self)
        layout.addWidget(QLabel(_("Select CSV Delimiter:")))

        self.radio_comma = QRadioButton(_("Comma ( , )"))
        self.radio_semicolon = QRadioButton(_("Semicolon ( ; )"))
        self.radio_comma.setChecked(True)

        layout.addWidget(self.radio_comma)
        layout.addWidget(self.radio_semicolon)

        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def get_delimiter(self) -> str:
        return "," if self.radio_comma.isChecked() else ";"


def import_csv(parent_window):
    """Import CSV file and ADD/INSERT all records into the active database."""
    if not parent_window.db_adapter:
        return

    dialog = CsvDelimiterDialog(parent_window, title="Import CSV Delimiter")
    if dialog.exec() != QDialog.DialogCode.Accepted:
        return

    delimiter = dialog.get_delimiter()

    file_path, _filter = QFileDialog.getOpenFileName(
        parent_window, _("Import CSV File"), "", "CSV Files (*.csv);;All Files (*)"
    )
    if not file_path:
        return

    try:
        added_count = 0
        with open(file_path, "r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f, delimiter=delimiter)
            for row in reader:
                raw_key = row.get("key")
                if not raw_key or parent_window.db_adapter.key_exists(raw_key):
                    new_key = str(parent_window.db_adapter.get_max_key() + 1)
                else:
                    new_key = str(raw_key)

                stamp_data = (
                    new_key,
                    row.get("nbr", "0"),
                    row.get("year", ""),
                    row.get("type", "Poste"),
                    row.get("valuecolor", ""),
                    row.get("stampDescription1", ""),
                    row.get("stampDescription", ""),
                    row.get("width", "0"),
                    row.get("height", "0"),
                    row.get("sequence", "1000"),
                    row.get("ascii_seq", "1000"),
                    row.get("sub_nbr", "")
                )

                parent_window.db_adapter.insert_stamp(stamp_data)
                added_count += 1

        parent_window.is_dirty = True
        parent_window._refresh_table()
        QMessageBox.information(
            parent_window,
            _("Import Success"),
            _(f"Successfully added {added_count} stamp record(s) from CSV to the database.")
        )
    except Exception as e:
        QMessageBox.critical(parent_window, _("Import Error"), f"{_('Error importing CSV:')}\n{str(e)}")


def export_csv(parent_window):
    """Export current catalogue to CSV file."""
    if not parent_window.db_adapter:
        return

    dialog = CsvDelimiterDialog(parent_window, title="Export CSV Delimiter")
    if dialog.exec() != QDialog.DialogCode.Accepted:
        return

    delimiter = dialog.get_delimiter()

    default_filename = f"{parent_window.current_country}_Catalogue.csv"
    file_path, _filter = QFileDialog.getSaveFileName(
        parent_window, _("Export CSV File"), default_filename, "CSV Files (*.csv)"
    )
    if not file_path:
        return

    stamps = parent_window.db_adapter.get_stamps()

    try:
        with open(file_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f, delimiter=delimiter)
            writer.writerow(HEADERS)
            for row in stamps:
                writer.writerow(list(row))

        QMessageBox.information(parent_window, _("Export Success"), _("Catalogue exported to CSV successfully."))
    except Exception as e:
        QMessageBox.critical(parent_window, _("Export Error"), f"{_('Error exporting CSV:')}\n{str(e)}")


def import_xlsx(parent_window):
    """Import XLSX file and ADD/INSERT all records into the active database."""
    if not parent_window.db_adapter:
        return

    try:
        import openpyxl
    except ImportError:
        QMessageBox.critical(
            parent_window,
            _("Missing Dependency"),
            _("The 'openpyxl' package is required for Excel files.<br>Install it using: <code>pip install openpyxl</code>")
        )
        return

    file_path, _filter = QFileDialog.getOpenFileName(
        parent_window, _("Import Excel File"), "", "Excel Files (*.xlsx);;All Files (*)"
    )
    if not file_path:
        return

    try:
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))

        if not rows:
            QMessageBox.warning(parent_window, _("Import Error"), _("The selected Excel file is empty."))
            return

        headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
        header_map = {name: idx for idx, name in enumerate(headers)}

        added_count = 0
        for row in rows[1:]:
            if not any(row):
                continue

            def get_val(col_name, default=""):
                if col_name in header_map:
                    val = row[header_map[col_name]]
                    return str(val).strip() if val is not None else default
                return default

            raw_key = get_val("key", "")
            if not raw_key or parent_window.db_adapter.key_exists(raw_key):
                new_key = str(parent_window.db_adapter.get_max_key() + 1)
            else:
                new_key = str(raw_key)

            stamp_data = (
                new_key,
                get_val("nbr", "0"),
                get_val("year", ""),
                get_val("type", "Poste"),
                get_val("valuecolor", ""),
                get_val("stampDescription1", ""),
                get_val("stampDescription", ""),
                get_val("width", "0"),
                get_val("height", "0"),
                get_val("sequence", "1000"),
                get_val("ascii_seq", "1000"),
                get_val("sub_nbr", "")
            )

            parent_window.db_adapter.insert_stamp(stamp_data)
            added_count += 1

        parent_window.is_dirty = True
        parent_window._refresh_table()
        QMessageBox.information(
            parent_window,
            _("Import Success"),
            _(f"Successfully added {added_count} stamp record(s) from Excel to the database.")
        )
    except Exception as e:
        QMessageBox.critical(parent_window, _("Import Error"), f"{_('Error importing Excel file:')}\n{str(e)}")


def export_xlsx(parent_window):
    """Export current catalogue to XLSX file."""
    if not parent_window.db_adapter:
        return

    try:
        import openpyxl
    except ImportError:
        QMessageBox.critical(
            parent_window,
            _("Missing Dependency"),
            _("The 'openpyxl' package is required for Excel files.<br>Install it using: <code>pip install openpyxl</code>")
        )
        return

    default_filename = f"{parent_window.current_country}_Catalogue.xlsx"
    file_path, _filter = QFileDialog.getSaveFileName(
        parent_window, _("Export Excel File"), default_filename, "Excel Files (*.xlsx)"
    )
    if not file_path:
        return

    stamps = parent_window.db_adapter.get_stamps()

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = parent_window.current_country or "Catalogue"

        ws.append(HEADERS)
        for row in stamps:
            ws.append(list(row))

        wb.save(file_path)
        QMessageBox.information(parent_window, _("Export Success"), _("Catalogue exported to Excel successfully."))
    except Exception as e:
        QMessageBox.critical(parent_window, _("Export Error"), f"{_('Error exporting Excel file:')}\n{str(e)}")